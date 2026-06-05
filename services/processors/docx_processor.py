"""DOCX content extraction using python-docx."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from services.processors.base import ExtractedDocument

logger = logging.getLogger(__name__)


class DOCXProcessor:
    """Extracts content, headings, and core metadata from DOCX files."""

    supported_extensions = {".docx"}

    def process(self, file_path: Path | str) -> ExtractedDocument:
        """Extract a DOCX into the standard document schema."""

        path = Path(file_path)
        try:
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("python-docx is required for DOCX processing") from exc

            document = Document(path)
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            headings = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip() and paragraph.style and paragraph.style.name.startswith("Heading")
            ]
            return {
                "file_name": path.name,
                "headings": headings,
                "content": "\n".join(paragraphs),
                "metadata": self._extract_metadata(document),
                "file_type": "docx",
            }
        except Exception as exc:
            logger.exception("Failed to process DOCX %s", path)
            raise ValueError(f"Could not process DOCX file: {path}") from exc

    def _extract_metadata(self, document: Any) -> dict[str, Any]:
        properties = document.core_properties
        return {
            "author": properties.author,
            "category": properties.category,
            "comments": properties.comments,
            "content_status": properties.content_status,
            "created": properties.created.isoformat() if properties.created else None,
            "identifier": properties.identifier,
            "keywords": properties.keywords,
            "language": properties.language,
            "last_modified_by": properties.last_modified_by,
            "modified": properties.modified.isoformat() if properties.modified else None,
            "subject": properties.subject,
            "title": properties.title,
            "version": properties.version,
        }
